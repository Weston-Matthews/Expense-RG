# Generates fake, QuickBooks-style monthly expense reports so Expense-RG can be tried without QuickBooks.
# Every number here is random sample data.
#
# Usage:  python samples/make_sample_reports.py [output_folder] [current_year] [latest_month]
# Writes one workbook per location per year, e.g. "NORTH 2020.xlsx" and "NORTH 2019.xlsx".

import json
import os
import random
import sys

import openpyxl as xl

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ACCOUNTS = ["Payroll Expenses", "Rent Expense", "Utilities", "Fuel", "Repairs and Maintenance", "Insurance Expense", "Office Supplies"]


def write_report(path, year, months, rng):
    # Mimics the column layout Expense Extract produces: account names on the left,
    # one value column per month (every other column), then a TOTAL column; the last row holds totals.
    wb = xl.Workbook()
    ws = wb.active
    ws.title = "Report"
    first_col = 4
    col = first_col
    for m in range(months):
        ws.cell(row=1, column=col, value=MONTHS[m]+" "+str(year)[2:])
        col += 2
    ws.cell(row=1, column=col, value="TOTAL")
    ws.cell(row=2, column=2, value="Expense")
    row = 3
    for account in ACCOUNTS:
        ws.cell(row=row, column=3, value=account)
        col = first_col
        line_total = 0.0
        for m in range(months):
            value = round(rng.uniform(500, 15000), 2)
            ws.cell(row=row, column=col, value=value)
            line_total += value
            col += 2
        ws.cell(row=row, column=col, value=round(line_total, 2))
        row += 1
    ws.cell(row=row, column=2, value="Total Expense")
    col = first_col
    for m in range(months+1):
        values = [ws.cell(row=r, column=col).value for r in range(3, row)]
        ws.cell(row=row, column=col, value=round(sum(values), 2))
        col += 2
    wb.save(path)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "sample_reports"
    cur_year = int(sys.argv[2]) if len(sys.argv) > 2 else 2020
    latest_month = int(sys.argv[3]) if len(sys.argv) > 3 else 9
    here = os.path.dirname(os.path.abspath(__file__))
    with open(os.path.join(here, "..", "config.example.json"), encoding="utf-8") as f:
        locations = json.load(f)["locations"]
    os.makedirs(out, exist_ok=True)
    rng = random.Random(42)
    for loc in locations:
        write_report(os.path.join(out, loc["code"]+" "+str(cur_year)+".xlsx"), cur_year, latest_month, rng)
        write_report(os.path.join(out, loc["code"]+" "+str(cur_year-1)+".xlsx"), cur_year-1, 12, rng)
    print("Wrote sample reports to "+os.path.abspath(out))


if __name__ == "__main__":
    main()
