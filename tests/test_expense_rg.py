import json
import os

import openpyxl as xl

from conftest import ROOT


def test_example_config_loads(erg):
    locations = erg.load_locations(os.path.join(ROOT, "config.example.json"))
    assert len(locations) <= erg.MAX_LOCATIONS
    names = [name for name, code in locations]
    assert len(names) == len(set(names))


def test_too_many_locations_rejected(erg, tmp_path):
    cfg = tmp_path / "config.json"
    cfg.write_text(json.dumps({"locations": [{"name": "L"+str(i)} for i in range(13)]}))
    try:
        erg.load_locations(str(cfg))
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_code_defaults_to_name(erg, tmp_path):
    cfg = tmp_path / "config.json"
    cfg.write_text(json.dumps({"locations": [{"name": "Depot"}]}))
    assert erg.load_locations(str(cfg)) == [("Depot", "Depot")]


def _report(months):
    wb = xl.Workbook()
    ws = wb.active
    col = 4
    for i, m in enumerate(["Jan 20", "Feb 20", "Mar 20"][:months]):
        ws.cell(row=1, column=col, value=m)
        ws.cell(row=5, column=col, value=100.0*(i+1))
        col += 2
    ws.cell(row=1, column=col, value="TOTAL")
    return ws


def test_detect_latest_month(erg):
    assert erg.detect_latest_month(_report(3)) == 3
    assert erg.detect_latest_month(_report(1)) == 1


def test_find_month_total_reads_bottom_row(erg):
    ws = _report(3)
    assert erg.find_month_total(ws, "Feb") == 200.0
    assert erg.find_month_total(ws, "Dec") == "ERR"
